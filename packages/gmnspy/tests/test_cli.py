"""Tests for the GMNS-aware CLI (task 4.1b / issue #83).

Exercises the extended ``info`` (overrides datagrove's generic with a
GMNS-aware version) + the new ``quality`` command. Reuses datagrove's
:class:`CliRunner` infra; both ``--json`` and rich output paths
exercised.
"""

from __future__ import annotations

import json

from gmnspy.cli.app import app
from gmnspy.fixtures import leavenworth
from typer.testing import CliRunner

runner = CliRunner()


# ---------------------------------------------------------------------------
# info — GMNS-aware override
# ---------------------------------------------------------------------------


def test_gmns_info_json_includes_spec_version():
    """gmnspy info --json carries spec_version + link/node counts."""
    result = runner.invoke(app, ["info", "--json", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["spec_version"] == "0.97"
    assert isinstance(payload["links"], int) and payload["links"] > 0
    assert isinstance(payload["nodes"], int) and payload["nodes"] > 0


def test_gmns_info_respects_spec_override():
    """--spec 0.96 stamps a different spec_version on the output."""
    result = runner.invoke(app, ["info", "--json", "--spec", "0.96", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["spec_version"] == "0.96"


def test_gmns_info_tables_is_jq_friendly_list_of_objects():
    """tables[] entries must be objects with name+rows, not bare name strings.

    Regression for the v1.0 CLI walk-through finding: the original
    shape was ``tables: ["link", "node", ...]`` which broke the
    documented jq pipeline ``info --json | jq '.tables[] | {name, rows}'``
    with ``Cannot index string with string "name"``. The fix makes
    ``tables`` a list of objects.
    """
    result = runner.invoke(app, ["info", "--json", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    tables = payload["tables"]
    assert isinstance(tables, list) and tables
    for entry in tables:
        assert isinstance(entry, dict), f"expected dict, got {type(entry).__name__}: {entry!r}"
        assert "name" in entry and "rows" in entry
        assert isinstance(entry["name"], str)
        assert isinstance(entry["rows"], int)


def test_gmns_info_rich_runs():
    """Rich-mode info exits 0 and writes to stderr."""
    result = runner.invoke(app, ["info", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr


# ---------------------------------------------------------------------------
# quality
# ---------------------------------------------------------------------------


def test_quality_json_emits_issues_document():
    """gmnspy quality --json writes {header, issues: [...]} on stdout."""
    result = runner.invoke(app, ["quality", "--json", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    assert "issues" in payload
    # All emitted issues should be DATA_QUALITY category.
    assert all(i["category"] == "data_quality" for i in payload["issues"])
    # Leavenworth's residential streets at 40 mph fire the high-speed rule.
    codes = {i["code"] for i in payload["issues"]}
    assert "quality.high_speed_residential" in codes


def test_quality_never_exits_nonzero_on_warnings_only():
    """quality findings are WARNING/INFO — command exit stays 0 even with issues."""
    result = runner.invoke(app, ["quality", "--json", str(leavenworth.csv_dir())])
    assert result.exit_code == 0


# ---------------------------------------------------------------------------
# Help + generic commands still present
# ---------------------------------------------------------------------------


def test_app_help_lists_gmns_and_generic_commands():
    """`gmnspy --help` shows the generic (info, validate) + GMNS (quality) commands."""
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    assert "info" in result.stdout
    assert "validate" in result.stdout
    assert "quality" in result.stdout


def test_gmns_validate_loads_spec_for_csv_directory():
    """`gmnspy validate <csv-dir>` must auto-load the GMNS spec.

    Regression for the bug found during the v1.0 CLI walk-through on
    2026-05-26: ``gmnspy validate`` was inheriting the generic
    ``datagrove validate`` (no spec passed to ``Package.from_source``),
    so a CSV directory without a ``datapackage.json`` produced a
    vacuous "every real table is unexpected" report. The fix added a
    GMNS-aware override in :mod:`gmnspy.cli.commands.validate` that
    routes through :meth:`Network.from_source` so the spec actually
    loads.

    This test guards against re-introducing the bug: it runs the CLI
    against the bundled Leavenworth CSV fixture and asserts that none
    of the known GMNS core tables (``link``, ``node``, ``geometry``,
    ``lane``) are flagged as ``structural.unexpected_resource``.
    """
    result = runner.invoke(app, ["validate", "--json", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    assert "issues" in payload

    unexpected = [
        i
        for i in payload["issues"]
        if i.get("code") == "structural.unexpected_resource" and i.get("table") in {"link", "node", "geometry", "lane"}
    ]
    assert not unexpected, (
        "GMNS spec was not loaded — core tables flagged as unexpected. "
        f"Got: {[i.get('table') for i in unexpected]}. "
        "Did gmnspy.cli.commands.validate.register get removed?"
    )


def test_gmns_validate_writes_html_report(tmp_path):
    """`gmnspy validate --html <path>` writes the interactive map+table viewer.

    Since gmnspy 1.x the ``--html`` output is the gmnspy.reports
    network-on-map viewer (Leaflet map pane + filterable issue
    table), not the legacy datagrove table-only report. The viewer
    is composed via :func:`gmnspy.reports.render_validation_html`.
    """
    out = tmp_path / "report.html"
    result = runner.invoke(app, ["validate", "--html", str(out), str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    assert out.is_file()
    html = out.read_text(encoding="utf-8")
    assert html.startswith("<!DOCTYPE html") or "<html" in html[:200]
    # Map viewer markers — the gv-map div is the signature of the new renderer.
    assert 'id="gv-map"' in html
    # And Leaflet is inlined (no remote script src).
    assert "Leaflet" in html
    assert "<script src=" not in html


def test_gmns_validate_writes_csv_findings(tmp_path):
    """``gmnspy validate --csv <path>`` writes the findings as a flat CSV file."""
    out = tmp_path / "findings.csv"
    result = runner.invoke(app, ["validate", "--csv", str(out), str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    assert out.is_file()
    text = out.read_text(encoding="utf-8")
    # CSV header must include the standard columns the writer emits.
    assert "severity,category,code" in text.splitlines()[0]


def test_gmns_validate_writes_xlsx_findings(tmp_path):
    """``gmnspy validate --xlsx <path>`` writes the findings as an .xlsx workbook."""
    out = tmp_path / "findings.xlsx"
    result = runner.invoke(app, ["validate", "--xlsx", str(out), str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    assert out.is_file()
    # The file must be a real xlsx (zip-based); openpyxl can open it.
    import openpyxl

    wb = openpyxl.load_workbook(out, read_only=True)
    ws = wb.active
    header = next(ws.iter_rows(values_only=True))
    assert "severity" in header
    assert "code" in header


def test_gmns_validate_html_falls_back_when_reports_extra_missing(tmp_path, monkeypatch):
    """``--html`` must still write something when ``[reports]`` is missing.

    Simulates the case where jinja2 (or the gmnspy.reports package itself) is
    not importable. The command falls back to datagrove's table-only HTML
    and prints a stderr note pointing at the install command. Exit is 0
    (or whatever the validation verdict says); the fallback is operational,
    not an error.
    """
    out = tmp_path / "report.html"
    import sys

    # Drop the reports submodule from sys.modules so the CLI's lazy import
    # path takes the ImportError branch.
    monkeypatch.setitem(sys.modules, "gmnspy.reports", None)
    monkeypatch.setitem(sys.modules, "gmnspy.reports.html_map", None)

    result = runner.invoke(app, ["validate", "--html", str(out), str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    assert out.is_file()
    html = out.read_text(encoding="utf-8")
    # Fallback is datagrove's table-only report.
    assert html.startswith("<!DOCTYPE html") or "<html" in html[:200]
    # Stderr note steers the user toward the right install command.
    combined = (result.stderr or "") + (result.stdout or "")
    assert "[reports]" in combined or "gmnspy[reports]" in combined


def test_gmns_validate_respects_spec_override():
    """`gmnspy validate --spec 0.96` should load the 0.96 spec, not the default."""
    result = runner.invoke(app, ["validate", "--json", "--spec", "0.96", str(leavenworth.csv_dir())])
    assert result.exit_code == 0, result.stderr
    # spec_version isn't always in the payload header, but the
    # important thing is the run succeeds against 0.96 (a different
    # spec from the 0.97 default).
